import docx
import pytest
from helpers import make_pdf

from thesis_radar.extract import Extracted, ExtractionError, extract, html_to_text


def test_text_and_markdown(tmp_path):
    path = tmp_path / "note.md"
    path.write_text("\n# My ACME notes\n\nPricing held.\n", encoding="utf-8")
    result = extract(path)
    assert result.pages == ["\n# My ACME notes\n\nPricing held.\n"]
    assert result.title == "My ACME notes"
    assert result.has_text


def test_html_drops_scripts_and_keeps_paragraphs():
    markup = (
        "<html><head><title>x</title><script>alert(1)</script></head>"
        "<body><p>First &amp; best.</p><p>Second<br>line.</p></body></html>"
    )
    assert html_to_text(markup) == "First & best.\n\nSecond\nline."


def test_html_skips_inline_xbrl_headers_and_hidden_elements():
    markup = (
        '<html><body><div style="display:none"><ix:header><ix:hidden>'
        "<ix:nonNumeric>false</ix:nonNumeric><xbrli:context><xbrli:identifier>0000320193</xbrli:identifier>"
        "</xbrli:context></ix:hidden></ix:header><div>still hidden</div></div>"
        '<p>Revenue grew <ix:nonFraction name="us-gaap:Revenue">12</ix:nonFraction>%.</p>'
        "<div hidden>secret</div><p>Visible.</p></body></html>"
    )
    text = html_to_text(markup)
    assert text == "Revenue grew 12%.\n\nVisible."


def test_html_table_cells_are_separated():
    markup = "<table><tr><td>Revenue</td><td>$</td><td>1,234</td></tr><tr><th>Units</th><th>50</th></tr></table>"
    assert html_to_text(markup) == "Revenue | $ | 1,234\n\nUnits | 50"


def test_docx_with_table(tmp_path):
    document = docx.Document()
    document.add_paragraph("ACME expert call")
    document.add_paragraph("Dealers are discounting.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Units"
    table.rows[0].cells[1].text = "50"
    path = tmp_path / "call.docx"
    document.save(str(path))
    result = extract(path)
    assert result.pages == ["ACME expert call\n\nDealers are discounting.\n\nUnits | 50"]
    assert result.title == "ACME expert call"


def test_pdf_pages_and_metadata_title(tmp_path):
    path = tmp_path / "report.pdf"
    path.write_bytes(make_pdf(["Inventory rose sharply.", "Pricing held firm."], title="ACME Q3 Preview"))
    result = extract(path)
    assert [page.strip() for page in result.pages] == ["Inventory rose sharply.", "Pricing held firm."]
    assert result.title == "ACME Q3 Preview"


def test_junk_pdf_titles_fall_back_to_text(tmp_path):
    path = tmp_path / "report.pdf"
    path.write_bytes(make_pdf(["ACME dealer survey"], title="Microsoft Word - draft3.docx"))
    assert extract(path).title == "ACME dealer survey"


def test_pdf_without_text_has_no_text(tmp_path):
    path = tmp_path / "scan.pdf"
    path.write_bytes(make_pdf([""]))
    assert not extract(path).has_text


def test_unsupported_and_corrupt_files_raise(tmp_path):
    odd = tmp_path / "data.xyz"
    odd.write_text("hi", encoding="utf-8")
    with pytest.raises(ExtractionError, match="unsupported"):
        extract(odd)
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"not a pdf")
    with pytest.raises(ExtractionError, match="could not read"):
        extract(broken)


def test_hash_ignores_whitespace_differences():
    a = Extracted(pages=["Inventory  rose.\n\nPricing held."], title="t")
    b = Extracted(pages=["Inventory rose.", "Pricing held."], title="t")
    assert a.text_sha256() == b.text_sha256()
