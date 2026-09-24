"""Extract plain text, with page boundaries, from supported file types."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from html.parser import HTMLParser
from pathlib import Path

SUPPORTED_SUFFIXES = frozenset({".pdf", ".html", ".htm", ".docx", ".txt", ".md"})


class ExtractionError(Exception):
    """A file could not be turned into text."""


@dataclass(frozen=True)
class Extracted:
    pages: list[str]
    title: str

    @property
    def has_text(self) -> bool:
        return any(page.strip() for page in self.pages)

    def normalized(self) -> str:
        return " ".join(" ".join(self.pages).split())

    def text_sha256(self) -> str:
        return hashlib.sha256(self.normalized().encode("utf-8")).hexdigest()


def extract(path: Path) -> Extracted:
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ExtractionError(f"unsupported file type: {suffix or '(none)'}")
    try:
        if suffix == ".pdf":
            pages, meta_title = _pdf(path)
        elif suffix in {".html", ".htm"}:
            pages, meta_title = [html_to_text(path.read_text(encoding="utf-8", errors="replace"))], None
        elif suffix == ".docx":
            pages, meta_title = [_docx(path)], None
        else:
            pages, meta_title = [path.read_text(encoding="utf-8", errors="replace")], None
    except Exception as exc:
        raise ExtractionError(f"could not read {path.name}: {exc}") from exc
    return Extracted(pages=pages, title=_title(meta_title, pages, path))


_JUNK_TITLES = re.compile(r"^(microsoft (word|powerpoint) - |untitled|document\d*$)", re.I)


def _title(meta_title: str | None, pages: list[str], path: Path) -> str:
    if meta_title and meta_title.strip() and not _JUNK_TITLES.match(meta_title.strip()):
        return meta_title.strip()[:200]
    for page in pages:
        for line in page.splitlines():
            cleaned = line.strip().lstrip("#").strip()
            if cleaned:
                return cleaned[:200]
    return path.stem


def _pdf(path: Path) -> tuple[list[str], str | None]:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = [page.extract_text() or "" for page in reader.pages]
    metadata = reader.metadata
    return pages, (metadata.title if metadata is not None else None)


def _docx(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    blocks = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                blocks.append(" | ".join(cells))
    return "\n\n".join(blocks)


_BLOCK_TAGS = frozenset(
    {"p", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6", "table", "section", "article", "pre", "blockquote",
     "ul", "ol", "dl", "dt", "dd", "header", "footer", "caption"}
)
_CELL_TAGS = frozenset({"td", "th"})
# Elements whose text is never content. `ix:header` holds inline-XBRL contexts and hidden facts.
_SKIP_TAGS = frozenset({"script", "style", "head", "noscript", "template", "ix:header", "xbrli:context"})
_HIDDEN_STYLE = re.compile(r"display\s*:\s*none", re.I)


class _TextCollector(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._skip_tag: str | None = None
        self._skip_depth = 0

    def _hidden(self, tag: str, attrs: list[tuple[str, str | None]]) -> bool:
        if tag in _SKIP_TAGS:
            return True
        for name, value in attrs:
            if name == "style" and value and _HIDDEN_STYLE.search(value):
                return True
            if name == "hidden":
                return True
        return False

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if self._skip_tag is not None:
            if tag == self._skip_tag:
                self._skip_depth += 1
            return
        if self._hidden(tag, attrs):
            self._skip_tag, self._skip_depth = tag, 1
        elif tag == "br":
            self.parts.append("\n")
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n\n")

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if self._skip_tag is None and tag == "br":
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if self._skip_tag is not None:
            if tag == self._skip_tag:
                self._skip_depth -= 1
                if self._skip_depth == 0:
                    self._skip_tag = None
            return
        if tag in _CELL_TAGS:
            self.parts.append(" | ")
        elif tag in _BLOCK_TAGS:
            self.parts.append("\n\n")

    def handle_data(self, data: str) -> None:
        if self._skip_tag is None:
            self.parts.append(data)


def html_to_text(markup: str) -> str:
    collector = _TextCollector()
    collector.feed(markup)
    collector.close()
    text = "".join(collector.parts)
    text = text.replace("\xa0", " ")
    text = re.sub(r"[ \t\r\f\v]+", " ", text)
    text = re.sub(r"(?: ?\| ?)+(?=\n|$)", "", text)  # separators left at the end of a table row
    text = re.sub(r"(?<=\n)(?: ?\| ?)+", "", text)  # and at the start of one
    text = re.sub(r" *\n *", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()
